import fitz
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_chroma import Chroma
from langchain_core.documents import Document
from app.config import VECTOR_DB_DIR, CHUNK_SIZE, CHUNK_OVERLAP

class DocumentIngestor:

    def __init__(self, embedding_model):
        self.embedding_model = embedding_model

    def extract(self, file_path):
        pages = []
        if file_path.endswith(".pdf"):
            doc = fitz.open(file_path)
            for i, page in enumerate(doc):
                text = page.get_text()
                if text.strip():
                    pages.append((i+1, text))
        elif file_path.endswith(".docx"):
            doc = DocxDocument(file_path)
            full_text = "\n".join([p.text for p in doc.paragraphs])
            if full_text.strip():
                pages.append((1, full_text))
        else:
            raise ValueError("Unsupported file format")
        return pages

    def ingest(self, file_path):
        pages = self.extract(file_path)
        
        # تقسيم النصوص
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=CHUNK_SIZE,
            chunk_overlap=CHUNK_OVERLAP,
            separators=["\n\n", "\n", ".", " ", ""]
        )

        documents = []
        chunk_id = 0

        for page_num, text in pages:
            chunks = splitter.split_text(text)
            for chunk in chunks:
                documents.append(
                    Document(
                        page_content=chunk,
                        metadata={
                            "source": file_path,
                            "page": page_num,
                            "chunk_id": chunk_id
                        }
                    )
                )
                chunk_id += 1

        if documents:
            
            Chroma.from_documents(
                documents,
                self.embedding_model,
                persist_directory=VECTOR_DB_DIR
            )

        return len(documents)